import asyncio
import html
import io
import os
import sqlite3
import subprocess
import re

from datetime import date, datetime, time, timedelta
import threading
from typing import Tuple, List, Optional

from openpyxl import Workbook

from aiogram import Bot
from aiogram.types import Message
from aiogram_dialog import DialogManager, StartMode
from sqlalchemy.orm import joinedload, selectinload

from aiogram.types import BufferedInputFile, FSInputFile

from app.api.routes.alert import clear_queue_alert
from app.bot.consts.ethers import SCHEDULE
from app.bot.consts.other import AVERAGE_SONG_SWITCH_DELAY
from app.bot.models import Ether, Order
from app.bot.models.auto_moderation import AutoModeration
from app.bot.models.banned_user import BannedUser
from app.bot.models.block_phrase import BlockPhrase
from app.bot.models.day_state import DayState
from app.bot.models.volume_change_point import VolumeChangePoint
from app.bot.player.mpv_player import player, _get_order_play_source
from app.bot.repositories.uow import UnitOfWork
from app.bot.routers.order_menu import (
    extract_youtube_video_id,
    has_time_passed,
    search_song_by_url,
    ytmusic,
    remove_brackets,
)
from app.bot.services.feedback import get_user_message_id
from app.bot.services.song_downloader import (
    add_to_download_queue,
    delete_song,
    get_song_path,
    is_downloading,
)
from app.bot.services.volume_changer import VolumeChanger
from app.bot.states.alert_state import get_alert_state, set_alert_state
from app.bot.states.help import HelpStates
from app.bot.states.main import MainStates
from app.settings import settings


def get_text_after_command(message: Message) -> str | None:
    full_text = message.text
    if not full_text:
        return None

    command_end_index = full_text.find(" ")
    if command_end_index == -1:
        return None

    return full_text[command_end_index + 1 :]


def parse_dates_from_command(arg: str | None) -> list[date] | None:
    """
    Parses a single date (21.06), or a range (21.06-30.06), returns list of date objects.
    If arg is None, returns [today].
    Returns None if invalid.
    """
    if not arg or not arg.strip():
        return [datetime.now().date()]

    arg = arg.strip()
    date_pattern = r"^(\d{2})\.(\d{2})$"
    range_pattern = r"^(\d{2})\.(\d{2})-(\d{2})\.(\d{2})$"

    if re.match(date_pattern, arg):
        day, month = map(int, arg.split("."))
        try:
            d = date(datetime.now().year, month, day)
            return [d]
        except ValueError:
            return None

    m = re.match(range_pattern, arg)
    if m:
        day1, month1, day2, month2 = map(int, m.groups())
        try:
            start = date(datetime.now().year, month1, day1)
            end = date(datetime.now().year, month2, day2)
            if end < start:
                return None
            days = []
            cur = start
            while cur <= end:
                days.append(cur)
                cur += timedelta(days=1)
            return days
        except ValueError:
            return None

    return None


def parse_dates_and_reason(
    arg: str | None,
) -> Tuple[Optional[List[date]], Optional[str]]:
    """
    Parses a single date (21.06), a range (21.06-30.06), or just a reason.
    Returns (list of dates, reason). If no date is given, uses today.
    Returns (None, None) if invalid.
    """
    today = datetime.now().date()
    if not arg or not arg.strip():
        return [today], None

    lines = [line.strip() for line in arg.strip().splitlines() if line.strip()]
    if not lines:
        return [today], None

    first_line = lines[0]
    rest = "\n".join(lines[1:]).strip()
    reason = None

    m = re.match(r"^(\d{2})\.(\d{2})-(\d{2})\.(\d{2})(?:\s+(.*))?$", first_line)
    if m:
        day1, month1, day2, month2, first_reason = m.groups()
        try:
            start = date(today.year, int(month1), int(day1))
            end = date(today.year, int(month2), int(day2))
            if end < start:
                return None, None
            days = []
            cur = start
            while cur <= end:
                days.append(cur)
                cur += timedelta(days=1)
            reason = first_reason or ""
            if rest:
                reason = f"{reason.strip()}\n{rest}".strip()
            return days, reason if reason else None
        except ValueError:
            return None, None

    m = re.match(r"^(\d{2})\.(\d{2})(?:\s+(.*))?$", first_line)
    if m:
        day, month, first_reason = m.groups()
        try:
            d = date(today.year, int(month), int(day))
            reason = first_reason or ""
            if rest:
                reason = f"{reason.strip()}\n{rest}".strip()
            return [d], reason if reason else None
        except ValueError:
            return None, None

    return [today], arg.strip()


async def start(message: Message, dialog_manager: DialogManager):
    await dialog_manager.start(MainStates.main, mode=StartMode.RESET_STACK)


async def help_command(message: Message, dialog_manager: DialogManager):
    await dialog_manager.start(HelpStates.select, mode=StartMode.RESET_STACK)


async def skip(message, uow):
    today = datetime.now()
    order = await uow.orders.find_one(
        Ether.ether_date == today.date(),
        Ether.start_time <= today.time(),
        Ether.cancelled == False,
        Order.played == False,
        Order.confirmed == True,
        Order.play_start != None,
        options=[joinedload(Order.ether)],
        order=[Order.play_start.desc()],
    )

    if not order:
        return await message.answer("Зараз нічого не грає")

    next_order = await uow.orders.find_one(
        Order.ether_id == order.ether_id,
        Order.played == False,
        Order.confirmed == True,
        Order.play_start == None,
        order=[Order.decision_timestamp.asc()],
    )
    order.played = True

    player.stop_current()

    if next_order:
        next_order.play_start = datetime.now()
        source = await _get_order_play_source(next_order)
        if source:
            player.play(source)

    await uow.flush()
    await message.answer("Трек скіпнуто")


async def cancel(message: Message, bot: Bot, uow: UnitOfWork):
    reply_message = message.reply_to_message
    if reply_message is None:
        await message.reply(
            "Команда /cancel має бути реплаєм на повідомлення із замовленням"
        )
        return

    reply_message_id = reply_message.message_id
    order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
    if order is None:
        await message.reply(
            "Команда /cancel має бути реплаєм на повідомлення із замовленням"
        )
        return

    if order.played:
        await message.reply("Трек вже програв або вже пропущений")
        return

    order.played = True

    if order.play_start:
        next_order = await uow.orders.find_one(
            Order.ether_id == order.ether_id,
            Order.played == False,
            Order.confirmed == True,
            Order.play_start == None,
            order=[Order.decision_timestamp.asc()],
        )
        order.played = True

        player.stop_current()

        if next_order:
            next_order.play_start = datetime.now()
            source = await _get_order_play_source(next_order)
            if source:
                player.play(source)

    video_id = order.video_id
    is_same_song_orders_exists = await uow.orders.check_exists(
        Order.video_id == video_id,
        Order.played == False,
        Order.confirmed == True,
        Order.expected_play_time >= order.expected_play_time - timedelta(hours=1),
    )

    if not is_same_song_orders_exists and not is_downloading(video_id):
        delete_song(video_id)

    await uow.flush()
    await message.answer("Трек скасовано")

    reason = get_text_after_command(message)
    cancel_text = f"🚫 Трек було скасовано: {order.title}"
    if reason:
        cancel_text += "\nПричина: " + reason

    await bot.send_message(order.ordered_by, cancel_text)


async def now_playing(message: Message, uow: UnitOfWork):
    today = datetime.now()

    order = await uow.orders.find_one(
        Ether.ether_date == today.date(),
        Ether.start_time <= today.time(),
        Ether.cancelled == False,
        Order.played == False,
        Order.confirmed == True,
        Order.play_start != None,
        options=[joinedload(Order.ether)],
        order=[Order.play_start.desc()],
    )

    if not order:
        return await message.answer("Зараз нічого не грає")

    expected_time = (
        order.expected_play_time.strftime("%H:%M:%S")
        if order.expected_play_time
        else "—"
    )
    actual_time = order.play_start.strftime("%H:%M:%S") if order.play_start else "—"
    chat_formatted = str(settings.ADMINS_CHAT_ID).replace("-100", "")

    text = (
        f"{order.title}\n"
        f"https://t.me/c/{chat_formatted}/{settings.ADMINS_MODERATION_THREAD_ID}/{order.order_message_id}\n\n"
        f"⏰ Очікуваний старт: {expected_time}\n"
        f"▶️ Фактичний: {actual_time}"
    )

    await message.answer(text, parse_mode=None)


async def stop(message: Message, uow: UnitOfWork):
    today = datetime.now()
    ether = await uow.ethers.find_one(
        Ether.ether_date == today.date(),
        Ether.start_time <= today.time(),
        Ether.cancelled == False,
        Ether.end_time >= today.time(),
        options=[selectinload(Ether.orders)],
    )

    if ether is not None:
        for order in ether.orders:
            order.played = True

        await uow.flush()

    if ether is not None and ether.ether_date == today.date():
        player.stop()

    if ether is None:
        await message.answer("Етер не знайдено!")
        return

    await message.answer("Чергу зупинено")


async def stop_today(message: Message, uow: UnitOfWork):
    today = datetime.now()
    ethers = await uow.ethers.find(
        Ether.ether_date == today.date(),
        Ether.end_time >= today.time(),
        Ether.cancelled == False,
        options=[selectinload(Ether.orders)],
    )

    for ether in ethers:
        for order in ether.orders:
            order.played = True

    await uow.flush()

    player.stop()
    await message.answer("Чергу зупинено. Всі замовлення на сьогодні видалено!")


async def stop_all(message: Message, uow: UnitOfWork):
    today = datetime.now()
    today_ethers = await uow.ethers.find(
        Ether.ether_date == today.date(),
        Ether.end_time >= today.time(),
        Ether.cancelled == False,
        options=[selectinload(Ether.orders)],
    )

    for ether in today_ethers:
        for order in ether.orders:
            order.played = True

    next_days_ethers = await uow.ethers.find(
        Ether.ether_date > today.date(),
        Ether.cancelled == False,
        options=[selectinload(Ether.orders)],
    )

    for ether in next_days_ethers:
        for order in ether.orders:
            order.played = True

    await uow.flush()

    player.stop()
    await message.answer("Чергу зупинено. Всі замовлення видалено!")


async def force_play_song(uow: UnitOfWork, filename: str):
    today = datetime.now()
    order = await uow.orders.find_one(
        Ether.ether_date == today.date(),
        Ether.start_time <= today.time(),
        Ether.cancelled == False,
        Order.played == False,
        Order.confirmed == True,
        Order.play_start != None,
        options=[joinedload(Order.ether)],
        order=[Order.play_start.desc()],
    )

    if order:
        order.played = True
        await uow.flush()

    player.play(filename)


async def traktor(message: Message, uow: UnitOfWork):
    await force_play_song(uow, "music/traktor.mp3")
    await message.answer("Трактор їде митися!")


async def shark(message: Message, uow: UnitOfWork):
    await force_play_song(uow, "music/shark.mp3")
    await message.answer("Baby Shark Dance!")


async def snow(message: Message, uow: UnitOfWork):
    await force_play_song(uow, "music/snow.mp3")
    await message.answer("Сніжинки пушинки!")


async def holiday(message: Message, uow: UnitOfWork):
    arg = get_text_after_command(message)
    dates, _ = parse_dates_and_reason(arg)
    if not dates:
        await message.reply(
            "Невірний формат дати! Використовуйте /holiday або /holiday 21.06 або /holiday 21.06-30.06"
        )
        return

    for d in dates:
        if d.weekday() == 6:
            await message.answer(f"{d.strftime('%d.%m')} — неділя завжди вихідний день")
            continue

        current_state = await uow.day_state.find_one(DayState.state_date == d)
        if current_state:
            current_state.is_holiday = True
        else:
            await uow.day_state.create(DayState(state_date=d, is_holiday=True))

        ethers = await uow.ethers.find(
            Ether.ether_date == d,
            Ether.cancelled == False,
            options=[selectinload(Ether.orders)],
        )

        for ether in ethers:
            ether.cancelled = True
            for order in ether.orders:
                order.played = True

    await uow.flush()
    if datetime.now().date() in dates:
        player.stop()
    if len(dates) == 1:
        await message.answer(
            f"День {dates[0].strftime('%d.%m')} тепер вихідний! Минула черга на цей день очищена!"
        )
    else:
        await message.answer(
            f"Дні {dates[0].strftime('%d.%m')} — {dates[-1].strftime('%d.%m')} тепер вихідні! Минула черга на ці дні очищена!"
        )


async def unholiday(message: Message, uow: UnitOfWork):
    arg = get_text_after_command(message)
    dates, _ = parse_dates_and_reason(arg)
    if not dates:
        await message.reply(
            "Невірний формат дати! Використовуйте /unholiday або /unholiday 21.06 або /unholiday 21.06-30.06"
        )
        return

    for d in dates:
        current_state = await uow.day_state.find_one(DayState.state_date == d)
        if current_state is None or not current_state.is_holiday:
            await message.answer(f"{d.strftime('%d.%m')} не був позначений як вихідний")
            continue

        current_state.is_holiday = False

        ethers = await uow.ethers.find(
            Ether.ether_date == d,
            Ether.cancelled == False,
            options=[selectinload(Ether.orders)],
        )

        for ether in ethers:
            ether.cancelled = True
            for order in ether.orders:
                order.played = True

    await uow.flush()
    if datetime.now().date() in dates:
        player.stop()
    if len(dates) == 1:
        await message.answer(
            f"День {dates[0].strftime('%d.%m')} тепер не вихідний! Минула черга на цей день очищена!"
        )
    else:
        await message.answer(
            f"Дні {dates[0].strftime('%d.%m')} — {dates[-1].strftime('%d.%m')} тепер не вихідні! Минула черга на ці дні очищена!"
        )


async def close(message: Message, uow: UnitOfWork):
    arg = get_text_after_command(message)
    dates, reason = parse_dates_and_reason(arg)
    if not dates:
        await message.reply(
            "Невірний формат дати! Використовуйте /close, /close 21.06, /close 21.06-23.06 або /close 21.06 причина"
        )
        return

    for d in dates:
        current_state = await uow.day_state.find_one(DayState.state_date == d)
        if current_state:
            current_state.is_closed = True
            current_state.reason = reason
        else:
            await uow.day_state.create(
                DayState(state_date=d, is_closed=True, reason=reason)
            )

        ethers = await uow.ethers.find(
            Ether.ether_date == d,
            Ether.cancelled == False,
            options=[selectinload(Ether.orders)],
        )

        for ether in ethers:
            for order in ether.orders:
                order.played = True

    await uow.flush()
    if datetime.now().date() in dates:
        player.stop()

    if len(dates) == 1:
        await message.answer(
            f"День {dates[0].strftime('%d.%m')} закритий для замовлень! Минула черга на цей день видалена"
        )
    else:
        await message.answer(
            f"Дні {dates[0].strftime('%d.%m')} — {dates[-1].strftime('%d.%m')} закриті для замовлень! Минула черга на ці дні видалена"
        )


async def open(message: Message, uow: UnitOfWork):
    arg = get_text_after_command(message)
    dates, _ = parse_dates_and_reason(arg)
    if not dates:
        await message.reply(
            "Невірний формат дати! Використовуйте /open, /open 21.06 або /open 21.06-23.06"
        )
        return

    for d in dates:
        current_state = await uow.day_state.find_one(DayState.state_date == d)
        if current_state:
            current_state.is_closed = False
            current_state.reason = None
        else:
            await uow.day_state.create(
                DayState(state_date=d, is_closed=False, reason=None)
            )

    await uow.flush()

    if len(dates) == 1:
        await message.answer(
            f"День {dates[0].strftime('%d.%m')} відкритий до замовлень!"
        )
    else:
        await message.answer(
            f"Дні {dates[0].strftime('%d.%m')} — {dates[-1].strftime('%d.%m')} відкриті до замовлень!"
        )


async def alert(message: Message, bot: Bot, uow: UnitOfWork):
    is_alert = await get_alert_state()
    if is_alert:
        await message.reply("Наразі вже триває тривога!")
        return

    await set_alert_state(True)
    await clear_queue_alert(uow, bot)

    player.play("music/alert.mp3")
    player.set_temp_volume(100)

    await message.reply("Повітряна тривога увімкненна!")


async def stop_alert(message: Message, uow: UnitOfWork):
    is_alert = await get_alert_state()
    if not is_alert:
        await message.reply("Тривоги наразі немає!")
        return

    await set_alert_state(False)

    player.play("music/all_clear.mp3")
    player.set_temp_volume(100)

    await message.reply("Повітряна тривога вимкнена!")


async def ban_with_feedback(message: Message, bot: Bot, uow: UnitOfWork):
    reply_message = message.reply_to_message
    if reply_message is None:
        await message.reply(
            "Команда /ban_with_feedback має бути реплаєм на повідомлення із замовленням або на повідомлення фідбеку"
        )
        return

    reply_message_id = reply_message.message_id
    user_id, _ = await get_user_message_id(reply_message_id)
    if user_id is None:
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Команда /ban_with_feedback має бути реплаєм на повідомлення із замовленням або на повідомлення фідбеку"
            )
            return

        user_id = order.ordered_by

    banned_user = await uow.banned_users.find_one(
        BannedUser.user_id == user_id, BannedUser.is_deleted == False
    )

    if banned_user:
        if banned_user.banned_feedback:
            await message.reply(
                "Користувач вже заблокований без можливості зворотнього звʼязку!"
            )
            return

        banned_user.banned_feedback = True
        await uow.flush()

        await message.reply("Юзеру вимкнено можливість зворотнього звʼязку!")

        block_reason = f" Причина блокування: {reason}" if reason else ""
        await bot.send_message(
            user_id,
            f"Тобі вимкнено можливість зворотнього звʼязку 😳.{block_reason}",
        )
        return

    reason = get_text_after_command(message)
    await uow.banned_users.create(
        BannedUser(
            user_id=user_id,
            ban_message_id=message.message_id,
            banned_by=message.from_user.id,
            timestamp=datetime.now(),
            reason=reason,
            banned_feedback=True,
        )
    )
    await uow.flush()
    await message.reply(
        f"Користувач з id <code>{user_id}</code> заблокований без можливості зворотнього звʼязку!",
        parse_mode="HTML",
    )

    block_reason = f" Причина блокування: {reason}" if reason else ""
    await bot.send_message(
        user_id,
        f"🚫 Тебе забанили. Тепер більше не зможеш писати в зворотний зв'язок та замовляти треки 🙃.{block_reason}",
    )


async def ban(message: Message, bot: Bot, uow: UnitOfWork):
    reply_message = message.reply_to_message
    if reply_message is None:
        await message.reply(
            "Команда /ban має бути реплаєм на повідомлення із замовленням або на повідомлення фідбеку"
        )
        return

    reply_message_id = reply_message.message_id
    user_id, _ = await get_user_message_id(reply_message_id)
    if user_id is None:
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Команда /ban має бути реплаєм на повідомлення із замовленням або на повідомлення фідбеку"
            )
            return

        user_id = order.ordered_by

    is_banned = await uow.banned_users.check_exists(
        BannedUser.user_id == user_id, BannedUser.is_deleted == False
    )

    if is_banned:
        await message.reply("Користувач вже заблокований!")
        return

    reason = get_text_after_command(message)
    await uow.banned_users.create(
        BannedUser(
            user_id=user_id,
            ban_message_id=message.message_id,
            banned_by=message.from_user.id,
            timestamp=datetime.now(),
            reason=reason,
            banned_feedback=False,
        )
    )
    await uow.flush()
    await message.reply(
        f"Користувач з id <code>{user_id}</code> заблокований!", parse_mode="HTML"
    )

    block_reason = f" Причина блокування: {reason}" if reason else ""
    additional_info = "Ти більше не можеш замовляти пісні. За потреби напиши модераторам, скориставшись функцією зворотнього зв'язку."
    try:
        await bot.send_message(
            user_id, f"🚫 Тебе забанили!{block_reason}\n\n{additional_info}"
        )
    except Exception:
        pass


async def unban(message: Message, bot: Bot, uow: UnitOfWork):
    user_id = get_text_after_command(message)

    if not user_id:
        await message.reply("Невірний формат команди! /unban user_id")
        return

    try:
        user_id = int(user_id)
    except ValueError:
        await message.reply("Невірний формат команди! /unban user_id")
        return

    banned_user = await uow.banned_users.find_one(
        BannedUser.user_id == user_id, BannedUser.is_deleted == False
    )

    if not banned_user:
        await message.reply("Цей користувач не заблокований!")
        return

    banned_user.is_deleted = True
    await uow.flush()
    await message.reply(
        f"Користувач з id <code>{user_id}</code> розблокований!", parse_mode="HTML"
    )

    try:
        await bot.send_message(
            user_id, "🎉 Тебе розблокували! Тепер знову можеш замовляти пісні!"
        )
    except Exception:
        pass


async def ban_list(message: Message, uow: UnitOfWork):
    banned_users = await uow.banned_users.find(BannedUser.is_deleted == False)

    if not banned_users:
        await message.reply("Список заблокованих користувачів порожній")
        return

    ban_list_message = f"<b>Заблоковані користувачі ({len(banned_users)})</b>\n"

    chat_id = settings.ADMINS_CHAT_ID
    chat_id_formatted = str(chat_id)[4:] if chat_id < 0 else str(chat_id)

    for i, user in enumerate(banned_users, 1):
        ban_message_url = f"https://t.me/c/{chat_id_formatted}/{user.ban_message_id}"
        feedback_emoji = "🔇" if user.banned_feedback else "🔊"
        ban_list_message += f'\n{i}) {feedback_emoji} <code>{user.user_id}</code> - <a href="{ban_message_url}">{user.timestamp.strftime("%d.%m.%Y")}</a>'

    await message.reply(ban_list_message, parse_mode="HTML")


async def set_volume(message: Message):
    volume = get_text_after_command(message)

    if not volume:
        await message.reply("Невірний формат команди! /volume гучність_цілим_числом")
        return

    try:
        volume = int(volume)
    except ValueError:
        await message.reply("Невірний формат команди! /volume гучність_цілим_числом")
        return

    if not 0 <= volume <= 100:
        await message.reply("Гучність має бути в межах 0-100")
        return

    if player.volume == volume:
        await message.reply("Наразі вже встановлена така гучність")
        return

    player.set_volume(volume)

    await message.reply(f"Гучність успішно встановлена на {volume}%")


async def set_temp_volume(message: Message):
    volume = get_text_after_command(message)

    if not volume:
        await message.reply(
            "Невірний формат команди! /temp_volume гучність_цілим_числом"
        )
        return

    try:
        volume = int(volume)
    except ValueError:
        await message.reply(
            "Невірний формат команди! /temp_volume гучність_цілим_числом"
        )
        return

    if not 0 <= volume <= 100:
        await message.reply("Гучність має бути в межах 0-100")
        return

    if player.volume == volume:
        await message.reply("Наразі вже встановлена така гучність")
        return

    player.set_temp_volume(volume)

    await message.reply(f"Гучність для поточної пісні успішно встановлена на {volume}%")


async def send_orders(message: Message, uow: UnitOfWork):
    await uow.flush()

    conn = sqlite3.connect("radio.db")
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM orders ORDER BY id DESC LIMIT 100;")
    orders_columns = [desc[0] for desc in cursor.description]
    orders_rows = cursor.fetchall()

    cursor.execute("SELECT * FROM ethers ORDER BY id DESC LIMIT 100;")
    ethers_columns = [desc[0] for desc in cursor.description]
    ethers_rows = cursor.fetchall()

    conn.close()

    buffer = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)

    orders_ws = wb.create_sheet(title="orders")
    orders_ws.append(orders_columns)
    for row in orders_rows:
        orders_ws.append(row)

    orders_ws.freeze_panes = "A2"

    for col_idx, column in enumerate(orders_columns, 1):
        col_values = [
            str(row[col_idx - 1]) for row in orders_rows if row[col_idx - 1] is not None
        ]
        max_length = max([len(str(column))] + [len(val) for val in col_values])
        orders_ws.column_dimensions[
            orders_ws.cell(row=1, column=col_idx).column_letter
        ].width = (max_length + 2)

    ethers_ws = wb.create_sheet(title="ethers")
    ethers_ws.append(ethers_columns)
    for row in ethers_rows:
        ethers_ws.append(row)

    ethers_ws.freeze_panes = "A2"

    for col_idx, column in enumerate(ethers_columns, 1):
        col_values = [
            str(row[col_idx - 1]) for row in ethers_rows if row[col_idx - 1] is not None
        ]
        max_length = max([len(str(column))] + [len(val) for val in col_values])
        ethers_ws.column_dimensions[
            ethers_ws.cell(row=1, column=col_idx).column_letter
        ].width = (max_length + 2)

    wb.save(buffer)
    buffer.seek(0)

    await message.reply_document(
        document=BufferedInputFile(file=buffer.getvalue(), filename="orders.xlsx")
    )


async def auto_moderation_list(message: Message, uow: UnitOfWork):
    await uow.flush()

    conn = sqlite3.connect("radio.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        WITH stats AS (
            SELECT
                video_id,
                MAX(title) AS title,
                MAX(duration) AS duration,
                SUM(CASE WHEN confirmed = 1 THEN 1 ELSE 0 END) AS confirmed,
                SUM(CASE WHEN confirmed = 0 AND decided_by <> 0 THEN 1 ELSE 0 END) AS rejected,
                SUM(
                    CASE
                        WHEN confirmed = 1 THEN  1
                        WHEN confirmed = 0 AND decided_by <> 0 THEN -1
                        ELSE 0
                    END
                ) AS rating
            FROM orders
            GROUP BY video_id
        )
        SELECT
            video_id,
            title,
            duration,
            confirmed,
            rejected,
            rating
        FROM stats
        WHERE rating > 2 OR rating < -2
        ORDER BY ABS(rating) ASC;
        """
    )
    orders_columns = [desc[0] for desc in cursor.description]
    orders_rows = cursor.fetchall()

    cursor.execute(
        """
        SELECT video_id, confirm, set_by, timestamp
        FROM auto_moderation
        WHERE is_deleted = FALSE
        ORDER BY
        CASE
            WHEN confirm IS NULL THEN 1
            WHEN confirm = TRUE THEN 2
            WHEN confirm = FALSE THEN 3
            ELSE 4
        END;
        """
    )
    lists_columns = [desc[0] for desc in cursor.description]
    lists_rows = cursor.fetchall()

    conn.close()

    buffer = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)

    autoapproved_ws = wb.create_sheet(title="Approve")
    autorejected_ws = wb.create_sheet(title="Reject")
    lists_ws = wb.create_sheet(title="Custom")

    autoapproved_ws.append(orders_columns)
    autorejected_ws.append(orders_columns)
    lists_ws.append(lists_columns)

    for row in orders_rows:
        if row[5] > 2:
            autoapproved_ws.append(row)
        else:
            autorejected_ws.append(row)

    for row in lists_rows:
        lists_ws.append(row)

    autoapproved_ws.freeze_panes = "A2"
    autorejected_ws.freeze_panes = "A2"
    lists_ws.freeze_panes = "A2"

    for col_idx, column in enumerate(orders_columns, 1):
        col_values = [
            str(row[col_idx - 1]) for row in orders_rows if row[col_idx - 1] is not None
        ]
        max_length = max([len(str(column))] + [len(val) for val in col_values])

        autoapproved_ws.column_dimensions[
            autoapproved_ws.cell(row=1, column=col_idx).column_letter
        ].width = (max_length + 2)

        autorejected_ws.column_dimensions[
            autorejected_ws.cell(row=1, column=col_idx).column_letter
        ].width = (max_length + 2)

    for col_idx, column in enumerate(lists_columns, 1):
        col_values = [
            str(row[col_idx - 1]) for row in lists_rows if row[col_idx - 1] is not None
        ]
        max_length = max([len(str(column))] + [len(val) for val in col_values])

        lists_ws.column_dimensions[
            lists_ws.cell(row=1, column=col_idx).column_letter
        ].width = (max_length + 2)

    wb.save(buffer)
    buffer.seek(0)

    await message.reply_document(
        document=BufferedInputFile(
            file=buffer.getvalue(), filename="auto-moderation.xlsx"
        )
    )


async def send_database_task(message: Message):
    conn = sqlite3.connect("radio.db")
    cursor = conn.cursor()

    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = [row[0] for row in cursor.fetchall()]

    buffer = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)

    for table in tables:
        cursor.execute(f"SELECT * FROM {table}")
        columns = [desc[0] for desc in cursor.description]
        rows = cursor.fetchall()

        ws = wb.create_sheet(title=table)
        ws.append(columns)

        for row in rows:
            ws.append(row)

        ws.freeze_panes = "A2"

        for col_idx, column in enumerate(columns, 1):
            col_values = [
                str(row[col_idx - 1]) for row in rows if row[col_idx - 1] is not None
            ]
            max_length = max([len(str(column))] + [len(val) for val in col_values])
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = (
                max_length + 2
            )

    conn.close()

    wb.save(buffer)
    buffer.seek(0)

    bot = Bot(settings.TOKEN.get_secret_value())
    await bot.send_document(
        document=BufferedInputFile(file=buffer.getvalue(), filename="database.xlsx"),
        chat_id=settings.ADMINS_CHAT_ID,
        reply_to_message_id=message.message_id,
        message_thread_id=message.message_thread_id,
    )
    await bot.session.close()


def run_in_thread_send_database(message: Message):
    asyncio.run(send_database_task(message))


async def send_database(message: Message, uow: UnitOfWork):
    await uow.flush()
    threading.Thread(
        target=run_in_thread_send_database, args=(message,), daemon=True
    ).start()


async def send_database_sql(message: Message, uow: UnitOfWork):
    await uow.flush()
    await message.reply_document(document=FSInputFile("radio.db"))


async def blacklist(message: Message, uow: UnitOfWork):
    video_id = get_text_after_command(message)

    if not video_id:
        reply_message = message.reply_to_message
        if reply_message is None:
            await message.reply(
                "Невірний формат команди! /blacklist video_id або реплай на замовлення"
            )
            return

        reply_message_id = reply_message.message_id
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Невірний формат команди! /blacklist video_id або реплай на замовлення"
            )
            return

        video_id = order.video_id

    if len(video_id) != 11:
        video_id = extract_youtube_video_id(video_id)
        if video_id is None:
            await message.reply("Невірне id відео або посилання!")
            return

    record = await uow.auto_moderation.find_one(
        AutoModeration.video_id == video_id, AutoModeration.is_deleted == False
    )

    if record and record.confirm == False:
        await message.reply("Це відео вже в blacklist!")
        return

    if record:
        record.is_deleted = True

    await uow.auto_moderation.create(
        AutoModeration(
            video_id=video_id,
            confirm=False,
            set_by=message.from_user.id,
            timestamp=datetime.now(),
        )
    )
    await uow.flush()

    await message.reply(
        f"Відео з id <code>{video_id}</code> внесено в blacklist!", parse_mode="HTML"
    )


async def whitelist(message: Message, uow: UnitOfWork):
    video_id = get_text_after_command(message)

    if not video_id:
        reply_message = message.reply_to_message
        if reply_message is None:
            await message.reply(
                "Невірний формат команди! /whitelist video_id або реплай на замовлення"
            )
            return

        reply_message_id = reply_message.message_id
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Невірний формат команди! /whitelist video_id або реплай на замовлення"
            )
            return

        video_id = order.video_id

    if len(video_id) != 11:
        video_id = extract_youtube_video_id(video_id)
        if video_id is None:
            await message.reply("Невірне id відео або посилання!")
            return

    record = await uow.auto_moderation.find_one(
        AutoModeration.video_id == video_id, AutoModeration.is_deleted == False
    )

    if record and record.confirm == True:
        await message.reply("Це відео вже в whitelist!")
        return

    if record:
        record.is_deleted = True

    await uow.auto_moderation.create(
        AutoModeration(
            video_id=video_id,
            confirm=True,
            set_by=message.from_user.id,
            timestamp=datetime.now(),
        )
    )
    await uow.flush()

    await message.reply(
        f"Відео з id <code>{video_id}</code> внесено в whitelist!", parse_mode="HTML"
    )


async def manual_list(message: Message, uow: UnitOfWork):
    video_id = get_text_after_command(message)

    if not video_id:
        reply_message = message.reply_to_message
        if reply_message is None:
            await message.reply(
                "Невірний формат команди! /manual_list video_id або реплай на замовлення"
            )
            return

        reply_message_id = reply_message.message_id
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Невірний формат команди! /manual_list video_id або реплай на замовлення"
            )
            return

        video_id = order.video_id

    if len(video_id) != 11:
        video_id = extract_youtube_video_id(video_id)
        if video_id is None:
            await message.reply("Невірне id відео або посилання!")
            return

    record = await uow.auto_moderation.find_one(
        AutoModeration.video_id == video_id, AutoModeration.is_deleted == False
    )

    if record and record.confirm == None:
        await message.reply("Це відео вже в manual list!")
        return

    if record:
        record.is_deleted = True

    await uow.auto_moderation.create(
        AutoModeration(
            video_id=video_id,
            confirm=None,
            set_by=message.from_user.id,
            timestamp=datetime.now(),
        )
    )
    await uow.flush()

    await message.reply(
        f"Відео з id <code>{video_id}</code> внесено в manual list!", parse_mode="HTML"
    )


async def remove_lists(message: Message, uow: UnitOfWork):
    video_id = get_text_after_command(message)

    if not video_id:
        reply_message = message.reply_to_message
        if reply_message is None:
            await message.reply(
                "Невірний формат команди! /remove_lists video_id або реплай на замовлення"
            )
            return

        reply_message_id = reply_message.message_id
        order = await uow.orders.find_one(Order.order_message_id == reply_message_id)
        if order is None:
            await message.reply(
                "Невірний формат команди! /remove_lists video_id або реплай на замовлення"
            )
            return

        video_id = order.video_id

    if len(video_id) != 11:
        video_id = extract_youtube_video_id(video_id)
        if video_id is None:
            await message.reply("Невірне id відео або посилання!")
            return

    record = await uow.auto_moderation.find_one(
        AutoModeration.video_id == video_id, AutoModeration.is_deleted == False
    )

    if not record:
        await message.reply("Для цього відео не задано жодних варіантів автомодерації!")
        return

    record.is_deleted = True

    await uow.flush()

    await message.reply(
        f"Відео з id <code>{video_id}</code> видалено зі списків автомодерації!",
        parse_mode="HTML",
    )


async def restart(message: Message, uow: UnitOfWork):
    try:
        await uow.flush()
    except Exception:
        pass

    script_path = os.path.abspath("./restart_bot.sh")
    subprocess.Popen(
        ["nohup", "bash", script_path],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    print("Bot restart script triggered. Exiting current instance.")
    await message.reply("🔄 Бот зараз перезапуститься!")


async def force_play(message: Message, uow: UnitOfWork):
    url = get_text_after_command(message)
    if not url:
        await message.reply("Невірний формат команди! /force_play {url}")
        return

    today = datetime.now()
    order = await uow.orders.find_one(
        Ether.ether_date == today.date(),
        Ether.start_time <= today.time(),
        Ether.cancelled == False,
        Order.played == False,
        Order.confirmed == True,
        Order.play_start != None,
        options=[joinedload(Order.ether)],
        order=[Order.play_start.desc()],
    )

    if order:
        order.played = True
        await uow.flush()

    video_info = await search_song_by_url(url, message, uow, False)
    if video_info is None:
        return

    video_id = video_info["video_id"]

    song_path = get_song_path(video_id)
    if song_path:
        player.play(str(song_path))
    else:
        player.play(f"https://youtube.com/watch?v={video_id}")

    await message.reply(f"⏯️ Примусово програється: {video_info['title']}")


async def force_play_playlist(message: Message, uow: UnitOfWork, is_once: bool):
    url = get_text_after_command(message)
    if not url:
        await message.reply(
            "Невірний формат команди! /команда {youtube_playlist_url}"
        )
        return

    now = datetime.now()
    ether = await uow.ethers.find_one(
        Ether.ether_date == now.date(),
        Ether.start_time <= now.time(),
        Ether.cancelled == False,
        Ether.end_time >= now.time(),
        options=[selectinload(Ether.orders)],
    )

    if ether is None:
        today = date.today()
        day_state = await uow.day_state.find_one(DayState.state_date == today)

        if day_state and day_state.is_holiday:
            day_schedule = SCHEDULE.get("6")
        else:
            day_schedule = SCHEDULE.get(str(now.weekday()))

        cur_time = now.time()

        if day_schedule is None:
            await message.answer("Етер не знайдено!")
            return

        for cur_ether in day_schedule:
            if has_time_passed(
                cur_time, cur_ether.get("start")
            ) and not has_time_passed(cur_time, cur_ether.get("end")):
                start_hour, start_minute = map(int, cur_ether["start"].split(":"))
                start_time = time(start_hour, start_minute)

                end_hour, end_minute = map(int, cur_ether["end"].split(":"))
                end_time = time(end_hour, end_minute)

                ether = await uow.ethers.create(
                    Ether(
                        name=cur_ether["name"],
                        start_time=start_time,
                        end_time=end_time,
                        ether_date=today,
                        cancelled=False,
                    )
                )

                break

    if ether is None:
        await message.answer("Етер не знайдено!")
        return

    try:
        playlist_id = None
        if "list=" in url:
            playlist_id = url.split("list=")[-1].split("&")[0]
        elif "/playlist/" in url:
            playlist_id = url.split("/playlist/")[-1].split("?")[0]
        if not playlist_id:
            await message.reply("Не вдалося визначити playlist_id з посилання.")
            return
        playlist = ytmusic.get_playlist(playlist_id, limit=None)
        tracks = playlist.get("tracks", [])
    except Exception as e:
        await message.reply(f"Помилка при отриманні плейлиста: {e}")
        return

    if not tracks:
        await message.reply("Плейлист порожній або не вдалося отримати треки.")
        return

    ether_start = max(now, datetime.combine(ether.ether_date, ether.start_time))
    ether_end = datetime.combine(ether.ether_date, ether.end_time)
    ether_duration = int((ether_end - ether_start).total_seconds())
    total_playlist_duration = 0

    for t in tracks:
        try:
            total_playlist_duration += int(t.get("duration_seconds", 0))
        except Exception:
            pass
    if total_playlist_duration == 0:
        await message.reply("Не вдалося визначити тривалість треків у плейлисті.")
        return

    for order in ether.orders:
        order.played = True

    await uow.flush()

    player.stop()

    orders = []
    seconds_filled = 0
    order_idx = 0
    admin_id = message.from_user.id

    while seconds_filled < ether_duration:
        for track in tracks:
            if seconds_filled >= ether_duration:
                break

            video_id = track.get("videoId")
            title = remove_brackets(track.get("title", ""))

            duration = int(track.get("duration_seconds", 0))
            if not video_id or not title or duration == 0:
                continue

            expected_play_time = now + timedelta(seconds=seconds_filled)
            order = Order(
                title=title,
                video_id=video_id,
                duration=duration,
                ether_id=ether.id,
                confirmed=True,
                ordered_by=admin_id,
                decision_timestamp=datetime.now(),
                played=False,
                play_start=datetime.now() if order_idx == 0 else None,
                expected_play_time=expected_play_time,
            )

            await uow.orders.create(order)

            orders.append(order)

            if order_idx == 0:
                song_path = get_song_path(video_id)
                if song_path:
                    player.play(str(song_path))
                else:
                    player.play(f"https://youtube.com/watch?v={video_id}")
            else:
                await add_to_download_queue(video_id)

            seconds_filled += duration + AVERAGE_SONG_SWITCH_DELAY
            order_idx += 1

        if is_once:
            break

    await uow.flush()

    if is_once:
        await message.reply(f"Поcтавлено на програвання {order_idx} треків з плейліста")
    else:
        await message.reply(
            f"Поточний етер заповнено треками з плейлиста! Всього додано: {order_idx} треків."
        )


async def force_playlist(message: Message, uow: UnitOfWork):
    await force_play_playlist(message, uow, False)


async def force_playlist_once(message: Message, uow: UnitOfWork):
    await force_play_playlist(message, uow, True)


async def add_volume_change_point(message: Message, uow: UnitOfWork):
    args = get_text_after_command(message)
    if not args:
        await message.answer("Використання: /add_volume_change_point hh:mm volume")
        return

    try:
        time_str, volume_str = args.split()
    except ValueError:
        await message.answer(
            "Неправильний формат. Використання: /add_volume_change_point hh:mm volume"
        )
        return

    time_match = re.match(r"^([0-1]?[0-9]|2[0-3]):([0-5][0-9])$", time_str)
    if not time_match:
        await message.answer("Неправильний формат часу. Має бути hh:mm")
        return

    hour, minute = map(int, time_match.groups())
    point_time = time(hour=hour, minute=minute)

    try:
        volume = int(volume_str)
        if not 0 <= volume <= 100:
            raise ValueError()
    except ValueError:
        await message.answer("Гучність має бути числом від 0 до 100")
        return

    existing_point = await uow.volume_change_points.find_one(
        VolumeChangePoint.time == point_time
    )
    if existing_point:
        await message.answer(f"Точка зміни гучності вже існує для часу {time_str}")
        return

    await uow.volume_change_points.create(
        VolumeChangePoint(time=point_time, volume=volume)
    )
    await uow.flush()

    await VolumeChanger.load_volume_points(uow)

    await message.answer(f"Точку зміни гучності додано: {time_str} -> {volume}%")


async def delete_volume_change_point(message: Message, uow: UnitOfWork):
    time_str = get_text_after_command(message)
    if not time_str:
        await message.answer("Використання: /delete_volume_change_point hh:mm")
        return

    time_match = re.match(r"^([0-1]?[0-9]|2[0-3]):([0-5][0-9])$", time_str.strip())
    if not time_match:
        await message.answer("Неправильний формат часу. Має бути hh:mm")
        return

    hour, minute = map(int, time_match.groups())
    point_time = time(hour=hour, minute=minute)

    point = await uow.volume_change_points.find_one(
        VolumeChangePoint.time == point_time
    )
    if not point:
        await message.answer(f"Точка зміни гучності не знайдена для часу {time_str}")
        return

    await uow.volume_change_points.delete(VolumeChangePoint.time == point_time)
    await uow.flush()

    await VolumeChanger.load_volume_points(uow)

    await message.answer(f"Точку зміни гучності видалено для часу {time_str}")


async def list_volume_change_points(message: Message, uow: UnitOfWork):
    points = await uow.volume_change_points.find(order=[VolumeChangePoint.time])

    if not points:
        await message.answer("Немає збережених точок зміни гучності")
        return

    lines = ["Список точок зміни гучності:"]
    for point in points:
        lines.append(f"{point.time.strftime('%H:%M')} -> {point.volume}%")

    await message.answer("\n".join(lines))


async def add_block_phrase(message: Message, uow: UnitOfWork):
    phrase = get_text_after_command(message)
    if not phrase:
        await message.answer("Використання: /add_block_phrase одне або кілька слів")
        return

    phrase = phrase.strip().lower()
    existing = await uow.block_phrases.find_one(BlockPhrase.phrase == phrase)

    if existing:
        await message.answer("Така фраза вже існує")
        return

    await uow.block_phrases.create(
        BlockPhrase(
            phrase=phrase, set_by=message.from_user.id, timestamp=datetime.now()
        )
    )
    await uow.flush()
    await message.answer("Фраза додана в базу даних")


async def delete_block_phrase(message: Message, uow: UnitOfWork):
    phrase = get_text_after_command(message)
    if not phrase:
        await message.answer("Використання: /delete_block_phrase одне або кілька слів")
        return

    phrase = phrase.strip().lower()
    existing = await uow.block_phrases.find_one(BlockPhrase.phrase == phrase)

    if not existing:
        await message.answer("Такої фрази не існує")
        return

    await uow.block_phrases.delete(BlockPhrase.phrase == phrase)
    await uow.flush()
    await message.answer("Фраза видалена")


async def list_block_phrases(message: Message, uow: UnitOfWork):
    phrases = await uow.block_phrases.find()

    if not phrases:
        await message.answer("Немає збережених фраз для блокування")
        return

    lines = ["Список фраз для блокування:"]
    for i, phrase in enumerate(phrases, 1):
        lines.append(f"{i}) {html.escape(phrase.phrase)}")

    await message.answer("\n".join(lines))


async def not_moderated(message: Message, uow: UnitOfWork):
    today = datetime.now()
    ethers_today = await uow.ethers.find(
        Ether.ether_date == today.date(),
        Ether.end_time >= today.time(),
        Ether.cancelled == False,
        options=[selectinload(Ether.orders)],
    )
    other_ethers = await uow.ethers.find(
        Ether.ether_date > today.date(),
        Ether.cancelled == False,
        options=[selectinload(Ether.orders)],
    )

    ethers = list(ethers_today) + list(other_ethers)
    message_ids = sorted(
        [
            order.order_message_id
            for ether in ethers
            for order in ether.orders
            if order.decided_by is None and order.order_message_id and not order.played
        ]
    )

    if len(message_ids) == 0:
        await message.answer("Усе промодеровано!")
        return

    answer_text = "Очікують на модерацію:"
    chat = str(settings.ADMINS_CHAT_ID).replace("-100", "")
    thread = str(settings.ADMINS_MODERATION_THREAD_ID)
    for i, message_id in enumerate(message_ids, 1):
        answer_text += f"\n{i}) https://t.me/c/{chat}/{thread}/{message_id}"

    await message.answer(answer_text)
